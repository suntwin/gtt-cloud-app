"""evening_pack.py — one-click evening export: PDF (for reading, app colours) + Excel (exact numbers).

Put this file at the repo root (next to gtt_process.py).

Sections (same order you decide in):
  0 Summary              market, session, regime, index trends, themes, breadth counts, scan breadth, note
  1 Leadership           sector table (all columns), healthiest first
  2 Breakouts by sector  one block per sector (≥ 3 breakouts in 20 sessions), every column incl. Saved? + Circuit %
  3 Top setups by sector tightest scanned stocks in the 10 sectors with most volume today (Market Themes tab)
  4 Saved breakouts      the whole watchlist with today's columns
  5 Tomorrow's list      split P1 / P2 / P3

build_pack(...) returns the sections; to_pdf() / to_xlsx() render them; render_pack_button() is the Streamlit UI.
"""
from __future__ import annotations

import io
import os
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd

import gtt_process as g

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "fonts")
MIN_SECTOR_BREAKOUTS = 3
NO_WRAP = {"date", "trigger_date", "last_trigger_date", "Symbol", "Keep", "Tier", "Rating", "Trend"}
MIN_W = {"Sector": 58, "Industry": 50, "Reason": 120, "Label": 50, "Note": 60, "Tags": 56, "Tag": 56, "source": 40, "Saved": 56, "Saved?": 52, "Patterns": 48}
TOP_SETUP_SECTORS = 10
TOP_SETUP_ROWS = 15

# key -> (label, number format or None for text, max chars in the PDF for text)
COLS = {
    "Keep": ("List", None, 4), "Tier": ("Tier", None, 4), "Patterns": ("Pattern", None, 18),
    "Label": ("Situation", None, 22), "Symbol": ("Symbol", None, 14), "Saved?": ("Saved?", None, 16),
    "Saved": ("Saved", None, 26), "Note": ("My note", None, 40), "Reason": ("Reason", None, 70),
    "Tag": ("Saved as", None, 24), "Tags": ("Tags", None, 26), "Rating": ("Rating", None, 4),
    "date": ("Broke out", None, 10), "chg": ("BO day %", "%.1f", 0), "vol": ("BO vol x", "%.1f", 0),
    "since": ("% since BO", "%+.1f", 0),
    "_chg_percentclose": ("Chg %", "%.1f", 0), "_vol_ratio": ("Vol x", "%.1f", 0),
    "_avgvol_mln": ("Avg vol", "%.0f", 0), "dvol": ("Day vol", "%.0f", 0), "Adr": ("ADR", "%.1f", 0),
    "_circuit": ("Circuit %", "%.0f", 0), "Ti65": ("Ti65", "%.2f", 0), "_nr4": ("NR4 %", "%.2f", 0),
    "_rel_tightness_today": ("Rel tight", "%.2f", 0), "_rel_wk_dist": ("ADRs from 10w", "%.1f", 0),
    "_10madist": ("10MA %", "%.1f", 0), "_20madist": ("20MA %", "%.1f", 0), "Avg_RS": ("Avg RS", "%.0f", 0),
    "Sector": ("Sector", None, 24), "Industry": ("Industry", None, 26),
    "trigger_date": ("Breakout day", None, 10), "last_trigger_date": ("Latest BO", None, 10),
    "trigger_close": ("BO close", "%.2f", 0), "prev_close": ("Fail level", "%.2f", 0),
    "vol_ratio": ("BO vol x", "%.1f", 0), "age": ("Age (days)", "%.0f", 0), "source": ("Source", None, 16),
    "Last 5": ("Last 5", "%.0f", 0), "Last 10": ("Last 10", "%.0f", 0), "Last 20": ("Last 20", "%.0f", 0),
    "Trend": ("Trend", None, 3), "Per 100 in scan": ("Per 100 in scan", "%.1f", 0),
    "Holding %": ("Holding %", "%.0f", 0), "Stocks": ("Stocks", "%.0f", 0),
    "Item": ("Item", None, 30), "Value": ("Value", None, 160),
}


def _live_cols(mcfg):
    return (["_chg_percentclose", "_vol_ratio", "_avgvol_mln", "dvol", "Adr"]
            + (["_circuit"] if mcfg.get("market") == "NSE" else [])
            + ["Ti65", "_nr4", "_rel_tightness_today", "_rel_wk_dist", "_10madist", "_20madist"])


# ═════════════════════════════════════════════════════════════════ colours
def _parse_css(css):
    """'background-color:#x;color:#y;font-weight:bold' -> (bg, fg, bold)"""
    if not css:
        return None
    d = {}
    for part in str(css).split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            d[k.strip()] = v.strip()
    bg = d.get("background-color")
    fg = d.get("color")
    bold = d.get("font-weight") == "bold"
    if not (bg or fg or bold):
        return None
    return (bg, fg, bold)


def _lead_c(v, good):
    v = g._num(v)
    if not np.isfinite(v):
        return None
    if v >= good[0]:
        return ("#28a745", "white", True)
    if v >= good[1]:
        return ("#d4edda", "black", False)
    return None


CELL_FN = {
    "_chg_percentclose": g._c_chg, "chg": g._c_chg, "_vol_ratio": g._c_vol, "vol": g._c_vol,
    "_rel_wk_dist": g._c_relwk, "_10madist": g._c_madist, "_20madist": g._c_madist,
    "_rel_tightness_today": g._c_tight, "_circuit": g._c_circuit, "Rating": g._c_rating,
}


def table_styles(df, mcfg):
    """{(row_idx, col_key): (bg, fg, bold)} for one table, same rules as the app."""
    out = {}
    cols = list(df.columns)
    for c in cols:
        vals = df[c].tolist()
        if c in CELL_FN:
            for i, v in enumerate(vals):
                s = _parse_css(CELL_FN[c](v))
                if s:
                    out[(i, c)] = s
        elif c == "_avgvol_mln":
            floor = float(g.TOMORROW_DEFAULTS["min_liq"])
            for i, css in enumerate(g._liq_styles(df[c], floor)):
                s = _parse_css(css)
                if s:
                    out[(i, c)] = s
        elif c == "since":
            for i, v in enumerate(vals):
                v = g._num(v)
                if np.isfinite(v):
                    out[(i, c)] = (None, "#28a745" if v > 0 else "#dc3545", True)
        elif c in ("Saved?", "Saved", "Keep"):
            for i, v in enumerate(vals):
                if str(v).startswith("✓"):
                    out[(i, c)] = (None, "#28a745", True)
                elif v:
                    out[(i, c)] = (None, "#8a94a6", False)
        elif c == "Label":
            for i, v in enumerate(vals):
                s = _parse_css(g.SITUATION_COLOURS.get(str(v), ""))
                if s:
                    out[(i, c)] = s
        elif c == "Last 10":
            for i, v in enumerate(vals):
                s = _lead_c(v, (5, 3))
                if s:
                    out[(i, c)] = s
        elif c == "Holding %":
            for i, v in enumerate(vals):
                s = _lead_c(v, (70, 55))
                if s:
                    out[(i, c)] = s
        elif c == "Trend":
            for i, v in enumerate(vals):
                if v in ("↑", "↓"):
                    out[(i, c)] = (None, "#28a745" if v == "↑" else "#dc3545", True)
        elif c == "Reason":
            for i, v in enumerate(vals):
                if str(v).startswith("⚡"):
                    out[(i, c)] = ("#ffd8a8", "#7a3e00", True)
    return out


def fmt_cell(key, v):
    lab, fmt, _ = COLS.get(key, (key, None, 30))
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if fmt:
        x = g._num(v)
        return "" if not np.isfinite(x) else fmt % x
    if isinstance(v, bool):
        return "✓" if v else ""
    if isinstance(v, (list, tuple)):
        return " + ".join(str(x) for x in v)
    s = str(v)
    return "" if s in ("nan", "None", "NaT") else s


# ═════════════════════════════════════════════════════════════════ data
def _live_frame(base_df):
    if base_df is None or base_df.empty or "Symbol" not in base_df.columns:
        return None
    d = g.add_derived(base_df).copy()
    d["Symbol"] = d["Symbol"].astype(str).str.strip().str.upper()
    return d.drop_duplicates("Symbol").set_index("Symbol", drop=False)


def _add_live(df, live, mcfg, key="Symbol"):
    if live is None:
        return df
    k = df[key].astype(str).str.upper()
    for c in _live_cols(mcfg):
        if c in live.columns:
            df[c] = pd.to_numeric(k.map(live[c]), errors="coerce").values
    return df


def _sort_quality(t):
    """Same order as the app: healthy groups (holding ≥ 60%) first, then most breakouts in the last 10 sessions."""
    t = t.copy()
    t["_weak"] = pd.to_numeric(t.get("Holding %"), errors="coerce").lt(60)
    return t.sort_values(["_weak", "Last 10", "Holding %", "Last 20"], ascending=[True, False, False, False],
                         na_position="last").drop(columns="_weak")


def build_pack(sb, mcfg, base_df, tomorrow=None):
    """Returns (meta, sections). sections = [(title, [ (block_title, df, note) ])]."""
    m = mcfg["market"]
    live = _live_frame(base_df)
    sd = g.effective_session(base_df, mcfg, sb)[0] if base_df is not None and not base_df.empty else g.journal_session(mcfg)
    meta = {"market": m, "session": sd, "made": datetime.now()}
    sections = []

    # ── 0 Summary
    note = None
    try:
        note = g.load_market_note(sb, mcfg, g.journal_session(mcfg))
    except Exception:
        pass
    rows = [("Market", m), ("Session", str(sd)), ("Exported", meta["made"].strftime("%a %d %b %Y %H:%M"))]
    if note:
        f = note.get("fields") or {}
        rname, rlim = g.regime_limits(note)
        meta["regime"] = rname
        rows.append(("Regime", f"{rname} → max {rlim['max_gtts']} new GTTs, {rlim['risk']}" if rlim else (rname or "—")))
        idx = f.get("indices") or g.INDEX_NAMES.get(m, [])
        rows.append(("Index trends", " · ".join(f"{n}: {f.get(f'trend_{i}') or '—'}" for i, n in enumerate(idx))))
        rows.append(("Themes (note)", f.get("themes") or "—"))
        b = f.get("breadth") or {}
        try:
            b = {**b, **{k: v for k, v in (g.load_breadth_day(sb, mcfg, sd) or {}).items() if k in g.BREADTH_KEYS and v is not None}}
        except Exception:
            pass
        rows.append(("Breadth", " · ".join(f"{lab}: {b.get(k) if b.get(k) is not None else '–'}" for k, lab in g.BREADTH_COLS)))
        if note.get("auto"):
            rows.append(("Scan breadth", g._auto_line(note["auto"])))
        if note.get("kind") == "SKIP":
            rows.append(("Note", "Skipped: " + (note.get("body") or "")))
        else:
            rows.append(("Market note", (note.get("body") or "").strip()))
    else:
        rows.append(("Market note", "Not written for this session."))
    sections.append(("Summary", [("Market", pd.DataFrame(rows, columns=["Item", "Value"]), None)]))

    # ── shared: history, watchlist
    try:
        hist = g.load_breakout_history(sb, mcfg, days=45)
    except Exception:
        hist = pd.DataFrame()
    try:
        watch = g.load_active_watchlist(sb, mcfg)
    except Exception:
        watch = []
    wdf = pd.DataFrame(watch)
    saved_syms = set(wdf["symbol"].astype(str).str.upper()) if not wdf.empty else set()
    tags_by = {str(r["symbol"]).upper(): " + ".join(g.tags_of(r)) for r in watch}
    try:
        hist90 = g.load_breakout_history(sb, mcfg, days=90)
    except Exception:
        hist90 = hist
    sec_map, ind_map = g._saved_groups(wdf, live, hist90) if not wdf.empty else ({}, {})
    ever = set(hist.loc[hist["selected"].fillna(False).astype(bool), "Symbol"].astype(str).str.upper()) \
        if not hist.empty and "selected" in hist.columns else set()

    # ── 1 Leadership
    t, names = (pd.DataFrame(), {})
    if not hist.empty:
        t, names = g.sector_leadership(hist, base_df, "Sector")
    if not t.empty:
        t = _sort_quality(t)
        n_sess = hist["date"].nunique()
        sections.append(("Leadership over time", [(
            f"Sectors · {len(hist)} breakouts over {n_sess} sessions (newest {hist['date'].max()}) · healthy groups (holding ≥ 60%) by breakouts, weak ones last",
            t[[c for c in ["Sector", "Last 5", "Last 10", "Last 20", "Trend", "Per 100 in scan", "Holding %", "Stocks"] if c in t.columns]].reset_index(drop=True),
            "Holding % = breakouts still above their pre-breakout close. Trend needs 20 sessions of history to mean much.")]))

    # ── 2 Breakouts by sector
    blocks = []
    if not t.empty:
        for _, r in t.iterrows():
            sec = r["Sector"]
            if int(r.get("Last 20", 0)) < MIN_SECTOR_BREAKOUTS:
                continue
            lst = names.get(sec, pd.DataFrame(columns=["Symbol", "date", "chg", "vol", "since"])).copy()
            have = set(lst["Symbol"].astype(str).str.upper())
            # watchlist stocks of this sector with no recorded evening breakout
            extra = []
            for w in watch:
                s = str(w["symbol"]).upper()
                if s in have or sec_map.get(w["symbol"], sec_map.get(s)) != sec:
                    continue
                bc, pc = g._num(w.get("trigger_close")), g._num(w.get("prev_close"))
                now = g._num(live.at[s, "Last"]) if live is not None and s in live.index and "Last" in live.columns else np.nan
                extra.append({"Symbol": s, "date": str(w.get("last_trigger_date") or w.get("trigger_date") or "")[:10],
                              "chg": (bc / pc - 1) * 100 if np.isfinite(bc) and np.isfinite(pc) and pc else np.nan,
                              "vol": g._num(w.get("vol_ratio")),
                              "since": (now / bc - 1) * 100 if np.isfinite(bc) and bc and np.isfinite(now) else np.nan,
                              "_manual": True})
            if extra:
                lst = pd.concat([lst, pd.DataFrame(extra)], ignore_index=True)
            manual = set(lst.loc[lst["_manual"].eq(True).fillna(False).astype(bool), "Symbol"]) if "_manual" in lst.columns else set()

            def status(sym):
                s = str(sym).upper()
                if s in manual:
                    return "✓ Saved (manual)"
                if s in saved_syms:
                    return "✓ Saved"
                return "Removed" if s in ever else "Skipped"
            lst.insert(1, "Saved?", lst["Symbol"].map(status))
            lst = lst.drop(columns=[c for c in ["_manual"] if c in lst.columns])
            lst = _add_live(lst, live, mcfg)
            if "_rel_tightness_today" in lst.columns:
                lst = lst.sort_values("_rel_tightness_today", na_position="last")
            hold = r.get("Holding %")
            title = (f"{sec} · {int(r['Last 10'])} breakouts last 10 · holding "
                     f"{'' if pd.isna(hold) else int(hold)}% · trend {r.get('Trend', '')}")
            blocks.append((title, lst.reset_index(drop=True), None))
    if blocks:
        sections.append(("Breakouts by sector (tightest today first)", blocks))

    # ── 3 Top setups by sector
    blocks = []
    if live is not None and "Sector" in live.columns:
        d = live.reset_index(drop=True)
        hot = (d[(d["_vol_ratio"].fillna(0) >= 1.0) & d["Sector"].notna() & (d["Sector"] != "Unknown")]
               .groupby("Sector")["Symbol"].count().sort_values(ascending=False).head(TOP_SETUP_SECTORS))
        for sec, n_vol in hot.items():
            ss = d[d["Sector"] == sec].copy()
            n_all = len(ss)
            ss = ss.sort_values("_rel_tightness_today", na_position="last").head(TOP_SETUP_ROWS)
            ss["Saved"] = ss["Symbol"].map(lambda s: ("✓ " + (tags_by.get(s) or "saved")) if s in saved_syms else "")
            cols = ["Symbol", "Saved"] + [c for c in _live_cols(mcfg) if c in ss.columns] + (["Industry"] if "Industry" in ss.columns else [])
            blocks.append((f"{sec} · tightest {len(ss)} of {n_all} scanned · {int(n_vol)} with volume today",
                           ss[cols].reset_index(drop=True), None))
    if blocks:
        sections.append(("Top setups by sector (all scanned stocks, tightest first)", blocks))

    # ── 4 Saved breakouts
    if not wdf.empty:
        v = pd.DataFrame({
            "Symbol": wdf["symbol"].astype(str),
            "Sector": wdf["symbol"].map(sec_map), "Industry": wdf["symbol"].map(ind_map),
            "Tags": [" + ".join(g.tags_of(r)) for r in watch],
            "Rating": wdf["rating"].map(g.rating_label) if "rating" in wdf.columns else "—",
            "trigger_date": wdf.get("trigger_date"), "last_trigger_date": wdf.get("last_trigger_date"),
            "trigger_close": wdf.get("trigger_close"), "prev_close": wdf.get("prev_close"),
            "vol_ratio": wdf.get("vol_ratio"),
            "source": wdf["tags"].map(lambda x: ", ".join(x) if isinstance(x, list) else "") if "tags" in wdf.columns else "",
        })
        ref = v["last_trigger_date"].fillna(v["trigger_date"]) if "last_trigger_date" in v.columns else v["trigger_date"]
        v["age"] = ref.map(lambda s: (date.today() - pd.to_datetime(s).date()).days if isinstance(s, str) and s else np.nan)
        v = _add_live(v, live, mcfg)
        v["_r"] = pd.to_numeric(wdf.get("rating"), errors="coerce").fillna(0).values
        v = v.sort_values(["Sector", "_r", "Symbol"], ascending=[True, False, True], na_position="last").drop(columns="_r")
        sections.append(("Saved breakouts", [(f"{len(v)} stocks on the watchlist · grouped by sector, best rated first",
                                              v.reset_index(drop=True), None)]))

    # ── 5 Tomorrow's list
    view, src = None, ""
    if tomorrow and tomorrow.get("view") is not None and str(tomorrow.get("sd")) == str(sd):
        view, src = tomorrow["view"].copy(), "as shown in the Tomorrow's list panel (your ticks)"
    elif base_df is not None and not base_df.empty:
        try:
            _, ylist = g.load_last_list(sb, mcfg, before_date=sd)
            cfg = {**g.TOMORROW_DEFAULTS, **g.TIER_DEFAULTS.get(m, {}), "circuit_check": m == "NSE"}
            lab, _ = g.classify_tomorrow(base_df, set(ylist), watch, cfg, sd)
            view = lab[lab["Label"] != "SCANNED"].copy()
            view["Label"] = view["Label"].map(lambda x: g.LABEL_NAMES.get(x, x))
            src = "default rules (open Anticipation → Tomorrow's list before exporting for your exact list and ticks)"
        except Exception as e:
            src = f"could not build ({e})"
    if view is not None and not view.empty:
        if "Keep" in view.columns:
            view["Keep"] = view["Keep"].map(lambda x: "✓" if bool(x) and pd.notna(x) else "")
        view = _add_live(view, live, mcfg)
        cols = [c for c in ["Keep", "Patterns", "Label", "Symbol", "Reason", "Tag"] if c in view.columns] + \
               [c for c in _live_cols(mcfg) if c in view.columns] + [c for c in ["Avg_RS", "Sector"] if c in view.columns]
        blocks = []
        tiers = view["Tier"].fillna("").astype(str) if "Tier" in view.columns else pd.Series("", index=view.index)
        for tier, title in (("P1", "P1 · high ADR"), ("P2", "P2 · middle ADR"), ("P3", "P3 · low ADR"), ("", "No ADR")):
            part = view[tiers == tier].copy()
            if part.empty:
                continue
            part["_k"] = part["Keep"].eq("✓") if "Keep" in part.columns else False
            part = part.sort_values(["_k", "_rel_tightness_today"] if "_rel_tightness_today" in part.columns else ["_k"],
                                    ascending=[False, True] if "_rel_tightness_today" in part.columns else [False],
                                    na_position="last")
            blocks.append((f"{title} · {len(part)} stocks · {int(part['_k'].sum())} ticked", part[cols].reset_index(drop=True), None))
        if blocks:
            sections.append((f"Tomorrow's list · {src}", blocks))
    return meta, sections


# ═════════════════════════════════════════════════════════════════ PDF
def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    reg, bold = os.path.join(FONT_DIR, "DejaVuSansCondensed.ttf"), os.path.join(FONT_DIR, "DejaVuSansCondensed-Bold.ttf")
    if os.path.exists(reg) and os.path.exists(bold):
        if "DVC" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DVC", reg))
            pdfmetrics.registerFont(TTFont("DVC-B", bold))
        return "DVC", "DVC-B", True
    return "Helvetica", "Helvetica-Bold", False


_ASCII = {"✓": "Y", "★": "*", "↑": "up", "↓": "dn", "→": "-", "⚡": "!", "·": "-", "÷": "/", "≤": "<=", "≥": ">=", "—": "-", "–": "-"}


def to_pdf(meta, sections, mcfg):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from xml.sax.saxutils import escape

    reg, bold, uni = _fonts()
    fix = (lambda s: s) if uni else (lambda s: "".join(_ASCII.get(ch, ch) for ch in s))
    FS = 6.6
    page = landscape(A4)
    avail = page[0] - 20 * mm
    st_h1 = ParagraphStyle("h1", fontName=bold, fontSize=13, leading=16, spaceAfter=4, textColor=colors.HexColor("#111827"))
    st_h2 = ParagraphStyle("h2", fontName=bold, fontSize=8.5, leading=11, spaceBefore=6, spaceAfter=2,
                           textColor=colors.HexColor("#1f2937"))
    st_note = ParagraphStyle("n", fontName=reg, fontSize=6.5, leading=8, textColor=colors.HexColor("#6b7280"))
    st_cell = ParagraphStyle("c", fontName=reg, fontSize=FS, leading=FS + 1.6)
    st_cell_b = ParagraphStyle("cb", parent=st_cell, fontName=bold)

    def make_table(df):
        from reportlab.pdfbase.pdfmetrics import stringWidth
        keys = list(df.columns)
        styles = table_styles(df, mcfg)
        header = [fix(COLS.get(k, (k,))[0]) for k in keys]
        body = [[fix(fmt_cell(k, v)) for k, v in zip(keys, row)] for row in df.itertuples(index=False)]
        widths = []
        for j, k in enumerate(keys):
            maxc = COLS.get(k, (k, None, 30))[2] if COLS.get(k, (k, None, 30))[1] is None else 0
            vals = [r[j] for r in body]
            if maxc:
                vals = [v[:maxc] for v in vals]
            w = max([stringWidth(header[j], bold, FS)] + [stringWidth(v, reg, FS) for v in vals[:400]]) + 6
            widths.append(min(w, 260))
        total = sum(widths)
        if total > avail:      # squeeze the text columns first (they wrap), then everything
            text = [j for j, k in enumerate(keys) if COLS.get(k, (k, None))[1] is None and k not in NO_WRAP and widths[j] > 40]
            excess = total - avail
            tw = sum(widths[j] for j in text)
            for j in text:
                widths[j] = max(min(MIN_W.get(keys[j], 36), widths[j]), widths[j] - excess * widths[j] / tw) if tw else widths[j]
            total = sum(widths)
            if total > avail:
                widths = [w * avail / total for w in widths]
        wrap = {j for j, k in enumerate(keys) if COLS.get(k, (k, None))[1] is None and k not in NO_WRAP}
        data = [header]
        for i, row in enumerate(body):
            cells = []
            for j, v in enumerate(row):
                if j in wrap and stringWidth(v, reg, FS) > widths[j] - 5:
                    s = styles.get((i, keys[j]))
                    cells.append(Paragraph(escape(v), st_cell_b if s and s[2] else st_cell))
                else:
                    cells.append(v)
            data.append(cells)
        num_cols = [j for j, k in enumerate(keys) if COLS.get(k, (k, None))[1]]
        cmds = [("FONTNAME", (0, 0), (-1, 0), bold), ("FONTNAME", (0, 1), (-1, -1), reg),
                ("FONTSIZE", (0, 0), (-1, -1), FS), ("LEADING", (0, 0), (-1, -1), FS + 1.6),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5), ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f7f9")])]
        for j in num_cols:
            cmds.append(("ALIGN", (j, 1), (j, -1), "RIGHT"))
        for (i, k), (bg, fg, b) in styles.items():
            j = keys.index(k)
            if bg:
                cmds.append(("BACKGROUND", (j, i + 1), (j, i + 1), colors.HexColor(bg) if bg.startswith("#") else colors.toColor(bg)))
            if fg:
                cmds.append(("TEXTCOLOR", (j, i + 1), (j, i + 1), colors.HexColor(fg) if fg.startswith("#") else colors.toColor(fg)))
            if b:
                cmds.append(("FONTNAME", (j, i + 1), (j, i + 1), bold))
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle(cmds))
        return t

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(reg, 6.5)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(10 * mm, 6 * mm, fix(f"Evening pack · {meta['market']} · session {meta['session']} · "
                                               f"exported {meta['made']:%d %b %Y %H:%M}"))
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    story = []
    for si, (title, blocks) in enumerate(sections):
        if si:
            story.append(PageBreak())
        story.append(Paragraph(escape(fix(title)), st_h1))
        for bt, df, note in blocks:
            if df is None or df.empty:
                continue
            if list(df.columns) == ["Item", "Value"]:       # summary: wrap long values
                rows = [[Paragraph(escape(fix(str(a))), st_cell_b), Paragraph(escape(fix(str(b))).replace("\n", "<br/>"), st_cell)]
                        for a, b in df.itertuples(index=False)]
                tt = Table(rows, colWidths=[35 * mm, avail - 35 * mm], hAlign="LEFT")
                tt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f3f4f6"))]))
                story.append(tt)
                continue
            head = [Paragraph(escape(fix(bt)), st_h2)]
            if note:
                head.append(Paragraph(escape(fix(note)), st_note))
            tbl = make_table(df)
            story.append(KeepTogether(head + [tbl]) if len(df) <= 18 else head[0])
            if len(df) > 18:
                story.extend(head[1:] + [tbl])
            story.append(Spacer(1, 3 * mm))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
                            bottomMargin=12 * mm, title=f"Evening pack {meta['market']} {meta['session']}")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ═════════════════════════════════════════════════════════════════ Excel
def to_xlsx(meta, sections, mcfg):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    names = {"Summary": "Summary", "Leadership over time": "Leadership"}
    used = set()
    for title, blocks in sections:
        base = names.get(title) or title.split(" (")[0].split(" · ")[0]
        sn = base[:31]
        n = 2
        while sn in used:
            sn = f"{base[:28]} {n}"; n += 1
        used.add(sn)
        ws = wb.create_sheet(sn)
        ws.append([title]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)
        ws.append([f"{meta['market']} · session {meta['session']} · exported {meta['made']:%d %b %Y %H:%M}"])
        ws.cell(ws.max_row, 1).font = Font(italic=True, color="6B7280")
        widths = {}
        for bt, df, note in blocks:
            if df is None or df.empty:
                continue
            ws.append([])
            ws.append([bt]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
            if note:
                ws.append([note]); ws.cell(ws.max_row, 1).font = Font(italic=True, color="6B7280")
            keys = list(df.columns)
            ws.append([COLS.get(k, (k,))[0] for k in keys])
            hr = ws.max_row
            for j in range(1, len(keys) + 1):
                c = ws.cell(hr, j)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F2937")
            styles = table_styles(df, mcfg)
            for i, row in enumerate(df.itertuples(index=False)):
                out = []
                for k, v in zip(keys, row):
                    fmt = COLS.get(k, (k, None))[1]
                    if fmt:
                        x = g._num(v)
                        out.append(None if not np.isfinite(x) else float(x))
                    else:
                        out.append(fmt_cell(k, v) or None)
                ws.append(out)
                r = ws.max_row
                for j, k in enumerate(keys, start=1):
                    c = ws.cell(r, j)
                    fmt = COLS.get(k, (k, None))[1]
                    if fmt:
                        c.number_format = "+0.0" if fmt == "%+.1f" else ("0" if fmt.endswith(".0f") else
                                                                        "0.0" if fmt.endswith(".1f") else "0.00")
                    if k in ("Value", "Reason", "Note"):
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                    s = styles.get((i, k))
                    if s:
                        bg, fg, b = s
                        if bg and bg.startswith("#"):
                            c.fill = PatternFill("solid", fgColor=bg[1:].upper())
                        fgc = (fg[1:].upper() if fg and fg.startswith("#") else ("FFFFFF" if fg == "white" else "000000" if fg else None))
                        c.font = Font(bold=b, color=fgc) if fgc else Font(bold=b)
            for j, k in enumerate(keys, start=1):
                lab, fmt, maxc = COLS.get(k, (k, None, 30))
                lens = [len(fmt_cell(k, x)) for x in df[k].tolist()]
                q = int(np.percentile(lens, 90)) if lens else 8
                w = max(len(lab), 6) + 2 if fmt else min(max(len(lab), q) + 2, 90 if k == "Value" else 45)
                widths[j] = max(widths.get(j, 0), w)
        for j, w in widths.items():
            ws.column_dimensions[get_column_letter(j)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═════════════════════════════════════════════════════════════════ Streamlit
def render_pack_button(st, sb, mcfg, base_df):
    key = f"evening_pack_{mcfg['market']}"
    with st.container(border=True):
        c1, c2 = st.columns([1, 3])
        ready = base_df is not None and not base_df.empty and sb is not None
        go = c1.button("📄 Evening pack (PDF + Excel)", disabled=not ready, key=f"{key}_go",
                       help="Leadership, breakouts by sector, top setups by sector, saved breakouts and tomorrow's "
                            "list P1/P2/P3 — one PDF for you, one Excel with exact numbers for Claude.")
        c2.caption("Generate the scan first. For your exact Tomorrow's list (with your ticks), open Anticipation so the "
                   "list is built before you export." if not ready else
                   "Everything for tonight's review in two files. Send both (plus your open positions: symbol, entry, stop, "
                   "shares) to Claude for a ranked shortlist.")
        if go:
            with st.spinner("Building the evening pack…"):
                try:
                    meta, sections = build_pack(sb, mcfg, base_df, st.session_state.get(f"pack_tomorrow_{mcfg['market']}"))
                    st.session_state[key] = {"meta": meta, "pdf": to_pdf(meta, sections, mcfg),
                                             "xlsx": to_xlsx(meta, sections, mcfg),
                                             "summary": [(t, len(b)) for t, b in sections]}
                except Exception as e:
                    import traceback
                    st.error(f"Could not build the evening pack ({e}).")
                    with st.expander("Details (send this to Claude)"):
                        st.code(traceback.format_exc()[-3000:], language=None)
        pack = st.session_state.get(key)
        if pack:
            m, sd = pack["meta"]["market"], pack["meta"]["session"]
            d1, d2, d3 = st.columns([1, 1, 2])
            d1.download_button("⬇ PDF", pack["pdf"], f"evening_pack_{m}_{sd}.pdf", "application/pdf", key=f"{key}_pdf")
            d2.download_button("⬇ Excel", pack["xlsx"], f"evening_pack_{m}_{sd}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"{key}_xlsx")
            d3.caption(f"Built {pack['meta']['made']:%H:%M} · " + " · ".join(f"{t.split(' (')[0].split(' · ')[0]} ({n})"
                                                                         for t, n in pack["summary"]))
